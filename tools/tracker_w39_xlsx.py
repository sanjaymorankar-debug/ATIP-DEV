"""
Builds the owner's deployment workbook (W39) from docs/ATIP_MASTER_TRACKER.csv:
tools/tracker_w39.py calls build(). The owner's sheets keep their layout and get their
status columns refreshed (plus evidence columns); new sheets carry the whole feature list.
"""

from __future__ import annotations

import re
from collections import Counter, defaultdict

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HDR = PatternFill("solid", fgColor="1F3A5F")
FILLS = {"COMPLETED": "C6EFCE", "IMPLEMENTED BUT NOT VERIFIED": "DDEBF7", "IN PROGRESS": "FFEB9C",
         "BLOCKED": "F8CBAD", "YET TO START": "EDEDED", "N/A": "FFFFFF",
         "Completed": "C6EFCE", "In Progress": "FFEB9C", "Blocked": "F8CBAD", "Yet to Start": "EDEDED"}
DONE = ("COMPLETED", "IMPLEMENTED BUT NOT VERIFIED")

WAVE_TEXT = {
    "W21": "Data & technical analysis (VWAP, S/R, betas, multi-timeframe, sector breadth)",
    "W22": "Factor research platform (formulas, factors, approval gate, studies)",
    "W23": "Research & backtesting depth (optimisation, sensitivity, robustness)",
    "W24": "Machine learning (GBM / RF / ensemble, walk-forward validation, anomalies, decay)",
    "W25": "Portfolio risk (exposure, concentration, correlation, VaR / ES, optimiser, emergency exit)",
    "W26": "Dashboard signals table + stock history panel",
    "W27": "Data & scores (NSE fundamentals, ownership, F&O summary, live feed, pre-open)",
    "W28": "Strategy & AI (intraday scans, lists, strategy performance, news AI)",
    "W28b": "News weighting + announcement NLP",
    "W29": "Execution (stop orders, modify, reconciliation, broker health, slippage, live P&L)",
    "W30": "Advanced quant (value / size, IV, pairs short leg, events, microstructure)",
    "W31": "Production hardening (off-site backups, vault, API reference, breakers, rollback drill)",
    "W32": "Enterprise SaaS layer (tenants, onboarding, billing foundation, alerts, webhooks)",
    "W33": "QA / UAT preparation",
    "W34": "Execution microstructure (algos, impact, latency, event-driven OMS / backtester)",
    "W35": "Data platform (lake, ticks, depth, F&O chains, macro, MF / multi-asset, alt data)",
    "W36": "ML & strategy tooling (neural net, RL, AI assistant, strategy builder, derivatives factors)",
    "W37": "Brokers & multi-asset (vault, connectors, multi-broker import, paper options)",
    "W38": "Platform & compliance (PostgreSQL path, containers, compliance, privacy, PWA)",
    "W39": "Tracker reconciliation: PERF-001 detail, owner notes, feature-map gaps, retail parity",
}

GAP_MAP = {
    "Market intelligence": r"^(DP-|SC-0[89]|DB-0[1-5]|NS-)",
    "Technical analysis": r"^TA-",
    "Quant signals": r"^(AF-|QR-|SC-)",
    "Backtesting": r"^BT-",
    "Paper trading": r"^(EX-(01|02|08|09|10|13|18|20)|BR-05)$",
    "Active trading": r"^(EX-|RK-|BR-)",
    "Portfolio intelligence": r"^(PF-|SC-10|DB-1[27])",
    "Investor profiling": r"^INV-001$",
    "Goal planning": r"^GOL-001$",
    "Asset allocation": r"^(AAL-001|PF-1[045])$",
    "Portfolio rebalancing": r"^(RBL-001|PF-06)$",
    "Wealth aggregation": r"^(WLT-001|PF-12|DP-21)$",
    "Consumer UX": r"^(UX-01|DB-|ENT-09|INT-001)",
    "AI Investment Advisor": r"^(AIA-001|ML-16)$",
}

EXCLUSION_NOTES = {
    "NPS": "Not built (excluded).",
    "FDs": "Not built (excluded).",
    "Tax": "Not built: rebalancing is tax-neutral by design (wealth/rebalance.py); costs include STT / stamp / GST "
           "only as transaction charges.",
    "Insurance": "Not built (excluded).",
    "Mutual Funds": "Data only: AMFI daily NAVs are ingested (DP-21, the owner's note 'amfi for MF data'); no MF "
                    "planning / advice features.",
}


def _style_header(ws, row=1):
    for c in ws[row]:
        if c.value is not None:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = HDR
            c.alignment = Alignment(wrap_text=True, vertical="top")


def _fit(ws, widths=None, default=18, wrap_from=2):
    for i, col in enumerate(ws.iter_cols(min_row=1, max_row=1), 1):
        w = (widths or {}).get(i, default)
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=wrap_from):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")


def _paint(cell):
    f = FILLS.get(str(cell.value or "").split(" (")[0].strip())
    if f:
        cell.fill = PatternFill("solid", fgColor=f)


def _table(ws, header, data, widths=None, paint_cols=()):
    ws.append(header)
    for d in data:
        ws.append([d.get(h, "") if isinstance(d, dict) else d[i] for i, h in enumerate(header)])
    _style_header(ws)
    _fit(ws, widths)
    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}{ws.max_row}"
    for col in paint_cols:
        for r in range(2, ws.max_row + 1):
            _paint(ws.cell(r, col))


def _num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _wave_status(feats):
    if not feats:
        return "Yet to Start", "Yet to Start", "no features mapped"
    st = Counter(f["Status"] for f in feats)
    ts = Counter(f["Testing Status"] for f in feats)
    n = len(feats)
    done = sum(st[s] for s in DONE)
    blocked = [f["ID"] for f in feats if f["Status"] == "BLOCKED"]
    open_ = [f["ID"] for f in feats if f["Status"] in ("IN PROGRESS", "YET TO START")]
    dev = "Completed" if done == n else ("In Progress" if open_ or done else "Blocked")
    if dev != "Completed" and not open_ and blocked:
        dev = "Completed (blocked items)" if done else "Blocked"
    tested = ts["COMPLETED"] + ts["N/A"]
    test = "Completed (automated)" if tested >= 0.9 * (n - len(blocked)) and n > len(blocked) else \
        ("In Progress" if ts["COMPLETED"] or ts["IN PROGRESS"] else "Yet to Start")
    why = (f"{done}/{n} features developed ({st['COMPLETED']} COMPLETED, {st['IMPLEMENTED BUT NOT VERIFIED']} "
           f"awaiting independent QA); automated tests: {ts['COMPLETED']} covered, {ts['IN PROGRESS']} partly, "
           f"{ts['YET TO START']} none")
    if blocked:
        why += f"; BLOCKED: {', '.join(blocked)}"
    if open_:
        why += f"; open: {', '.join(open_)}"
    return dev, test, why


def build(owner_xlsx, out, fields, rows, diff):
    wb = load_workbook(owner_xlsx)
    by_wave = defaultdict(list)
    for r in rows:
        by_wave[r.get("Deployment Wave", "")].append(r)
    by_id = {r["ID"]: r for r in rows}

    # ── Master Tracker: refresh every wave row, append the repo waves W21-W39 ──
    ws = wb["Master Tracker"]
    hdr = [c.value for c in ws[1]]
    ci = {h: i + 1 for i, h in enumerate(hdr)}
    for r in range(2, ws.max_row + 1):
        wave = ws.cell(r, 1).value
        if wave is None:
            continue
        key = str(wave).rstrip("0").rstrip(".") if isinstance(wave, float) else str(wave)
        feats = by_wave.get(key, [])
        if key == "9":
            feats = feats + by_wave.get("9+ (Enterprise SaaS)", [])
        dev, test, why = _wave_status(feats)
        nxt = sorted({f["Blocker / Input Required"] for f in feats if f.get("Blocker / Input Required")})
        ws.cell(r, ci["Development Status"]).value = dev
        ws.cell(r, ci["Testing Status"]).value = test
        ws.cell(r, ci["Current Status / Rationale"]).value = f"[2026-10-07] {why}"
        ws.cell(r, ci["Next Action / Dependency"]).value = (" | ".join(nxt[:4]) if nxt else
                                                           "Independent QA (QA-001) + owner UAT (UAT-001)")
    for w, text in WAVE_TEXT.items():
        # "W25" but not the "W25-W32" of a range, nor the "W39 check" notes added to older rows
        tag = re.compile(rf"(?<![-\w]){w}\b(?!-)")
        if w == "W39":
            feats = [x for x in rows if x["Phase"] == "11" or "W39:" in x.get("Current Implementation", "")]
        else:
            feats = [x for x in rows if tag.search(x.get("Notes", "") + " " + x.get("Current Implementation", ""))]
        dev, test, why = _wave_status(feats)
        ws.append([w, text, "repository wave (platform extension)", "Platform extension", dev, test,
                   f"[2026-10-07] {why}", "Independent QA" if dev.startswith("Completed") else "see Blockers"])
    for r in range(2, ws.max_row + 1):
        for col in (ci["Development Status"], ci["Testing Status"]):
            _paint(ws.cell(r, col))
            ws.cell(r, col).alignment = Alignment(wrap_text=True, vertical="top")

    # ── PERF-001 Detail ──
    ws = wb["PERF-001 Detail"]
    hdr = [c.value for c in ws[1]]
    extra = ["Implementation (2026-10-07)", "Automated tests", "Remaining"]
    for j, h in enumerate(extra):
        ws.cell(1, len(hdr) + 1 + j).value = h
    ci = {h: i + 1 for i, h in enumerate(hdr + extra)}
    for r in range(2, ws.max_row + 1):
        f = by_id.get(str(ws.cell(r, 1).value))
        if not f:
            continue
        ws.cell(r, ci["Development Status"]).value = "Completed" if f["Status"] in DONE else f["Status"].title()
        ws.cell(r, ci["Testing Status"]).value = {"COMPLETED": "Completed (automated)", "IN PROGRESS": "In Progress",
                                                  "YET TO START": "Yet to Start"}.get(f["Testing Status"],
                                                                                      f["Testing Status"])
        ws.cell(r, ci[extra[0]]).value = f["Current Implementation"]
        ws.cell(r, ci[extra[1]]).value = f["Evidence"]
        ws.cell(r, ci[extra[2]]).value = f["Next Action / Missing Work"]
        for col in (ci["Development Status"], ci["Testing Status"]):
            _paint(ws.cell(r, col))
    _style_header(ws)

    # ── Gap Analysis: tracker-derived completion beside the owner's estimate ──
    ws = wb["Gap Analysis"]
    n0 = ws.max_column
    for j, h in enumerate(["Tracker-derived % (2026-10-07)", "Features counted", "Blocked in area"]):
        ws.cell(1, n0 + 1 + j).value = h
    for r in range(2, ws.max_row + 1):
        area = ws.cell(r, 1).value
        rx = GAP_MAP.get(area)
        if not rx:
            continue
        fs = [x for x in rows if re.match(rx, x["ID"])]
        pcts = [p for p in (_num(x["Completion %"]) for x in fs) if p is not None]
        ws.cell(r, n0 + 1).value = round(sum(pcts) / len(pcts), 1) if pcts else None
        ws.cell(r, n0 + 2).value = len(fs)
        ws.cell(r, n0 + 3).value = ", ".join(x["ID"] for x in fs if x["Status"] == "BLOCKED") or "-"
    _style_header(ws)

    # ── Scope Exclusions / Roadmap / Implementation Sequence: a status column each ──
    ws = wb["Scope Exclusions"]
    ws.cell(1, ws.max_column + 1).value = "Status in code (2026-10-07)"
    for r in range(2, ws.max_row + 1):
        ws.cell(r, ws.max_column).value = EXCLUSION_NOTES.get(ws.cell(r, 2).value, "")
    _style_header(ws)
    ws = wb["Implementation Sequence"]
    n0 = ws.max_column
    ws.cell(1, n0 + 1).value = "Development Status (2026-10-07)"
    ws.cell(1, n0 + 2).value = "Testing Status (2026-10-07)"
    for r in range(2, ws.max_row + 1):
        wave = ws.cell(r, 2).value
        key = str(wave).rstrip("0").rstrip(".") if isinstance(wave, float) else str(wave)
        feats = by_wave.get(key, []) + (by_wave.get("9+ (Enterprise SaaS)", []) if key == "9" else [])
        dev, test, _ = _wave_status(feats)
        ws.cell(r, n0 + 1).value, ws.cell(r, n0 + 2).value = dev, test
        _paint(ws.cell(r, n0 + 1))
        _paint(ws.cell(r, n0 + 2))
    _style_header(ws)
    ws = wb["Roadmap"]
    n0 = ws.max_column
    ws.cell(1, n0 + 1).value = "Status (2026-10-07)"
    for r in range(2, ws.max_row + 1):
        waves = re.findall(r"\d+(?:\.\d)?", str(ws.cell(r, 2).value or ""))
        if "–" in str(ws.cell(r, 2).value) and len(waves) == 2:
            waves = [str(i) for i in range(int(waves[0]), int(waves[1]) + 1)]
        if str(ws.cell(r, 3).value or "").startswith("Rebalancing + Performance"):
            waves.append("15.5")
        feats = [f for w in waves for f in by_wave.get(w, [])]
        ws.cell(r, n0 + 1).value = _wave_status(feats)[0]
        _paint(ws.cell(r, n0 + 1))
    _style_header(ws)

    # ── new sheets ──
    s = wb.create_sheet("Summary", 0)
    s.append(["ATIP feature tracker -- reconciled 2026-10-07 (W39)"])
    s["A1"].font = Font(bold=True, size=14)
    s.append(["Source", "docs/ATIP_MASTER_TRACKER.csv (canonical), the owner's uploaded workbook and CSV, the code "
                        "and a full automated test run with coverage"])
    s.append(["Features", len(rows)])
    s.append([])
    for title, key in (("Development status", "Status"), ("Testing status (automated)", "Testing Status")):
        s.append([title, "Count"])
        _hdr_row(s)
        for k, v in sorted(Counter(r[key] for r in rows).items(), key=lambda kv: -kv[1]):
            s.append([k, v])
            _paint(s.cell(s.max_row, 1))
        s.append([])
    s.append(["Deployment wave", "Features", "Developed", "Blocked", "Testing COMPLETED"])
    _hdr_row(s)
    for w in sorted(by_wave, key=_wave_sort):
        fs = by_wave[w]
        s.append([w or "(none)", len(fs), sum(1 for f in fs if f["Status"] in DONE),
                  sum(1 for f in fs if f["Status"] == "BLOCKED"),
                  sum(1 for f in fs if f["Testing Status"] == "COMPLETED")])
    s.append([])
    s.append(["Phase", "Features", "Developed", "Blocked"])
    _hdr_row(s)
    by_phase = defaultdict(list)
    for r in rows:
        by_phase[(int(r["Phase"]), r["Phase Name"])].append(r)
    for (p, name), fs in sorted(by_phase.items()):
        s.append([f"{p} {name}", len(fs), sum(1 for f in fs if f["Status"] in DONE),
                  sum(1 for f in fs if f["Status"] == "BLOCKED")])
    s.column_dimensions["A"].width = 48
    for c in "BCDE":
        s.column_dimensions[c].width = 16

    cols = fields
    widths = {i + 1: (12 if c in ("Seq", "Phase", "Priority", "Completion %", "Deployment Wave") else
                      44 if c in ("Current Implementation", "Next Action / Missing Work", "Notes", "Test Evidence",
                                  "Blocker / Input Required") else 22) for i, c in enumerate(cols)}
    paint = [cols.index("Status") + 1, cols.index("Testing Status") + 1]
    _table(wb.create_sheet("All Features", 1), cols, rows, widths, paint)
    new = [r for r in rows if r["Phase"] == "11"]
    _table(wb.create_sheet("New Items (W39)", 2), cols, new, widths, paint)
    blk = [r for r in rows if r["Status"] == "BLOCKED" or r.get("Blocker / Input Required")]
    bcols = ["ID", "Feature", "Status", "Completion %", "Blocker / Input Required", "Current Implementation",
             "Next Action / Missing Work"]
    _table(wb.create_sheet("Blockers", 3), bcols, blk, {1: 12, 2: 30, 3: 18, 4: 10, 5: 60, 6: 50, 7: 40}, [3])
    dcols = ["ID", "Feature", "Owner CSV status", "Owner CSV %", "Repo status (W39)", "Repo %", "Why"]
    _table(wb.create_sheet("Owner CSV vs repo"), dcols, diff, {1: 12, 2: 30, 3: 22, 4: 10, 5: 26, 6: 10, 7: 70},
           [3, 5])
    lg = wb.create_sheet("Status Legend")
    for line in (
            ["Development status (CSV vocabulary)", "Meaning"],
            ["COMPLETED", "In master, with tests / evidence; maintained"],
            ["IMPLEMENTED BUT NOT VERIFIED", "Developed and merged; automated tests where listed; independent QA "
                                             "(ChatGPT) and owner UAT still pending"],
            ["IN PROGRESS", "Partly built"],
            ["BLOCKED", "Cannot proceed without the input in 'Blocker / Input Required'"],
            ["YET TO START", "Not started"],
            [],
            ["Testing status (automated, development side)", "Meaning"],
            ["COMPLETED", ">= 60% of the key files' lines run by passing automated tests, or W39 tests written for "
                          "the feature"],
            ["IN PROGRESS", "1-59% of the key files' lines run by tests"],
            ["YET TO START", "No automated test reaches the key files"],
            ["BLOCKED / N/A", "Feature blocked / no code (documentation, legal, process)"],
            [],
            ["Owner workbook vocabulary", "Completed / In Progress / Blocked / Yet to Start; 'Completed (automated)' "
                                          "= automated tests pass, independent QA pending"]):
        lg.append(line)
    lg.column_dimensions["A"].width = 44
    lg.column_dimensions["B"].width = 100
    for c in (lg["A1"], lg["B1"], lg["A8"], lg["B8"], lg["A14"]):
        c.font = Font(bold=True)
    wb.save(out)


def _hdr_row(ws):
    for c in ws[ws.max_row]:
        if c.value is not None:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = HDR


def _wave_sort(w):
    m = re.match(r"(\d+(?:\.\d+)?)", w or "")
    return (float(m.group(1)) if m else 999, w)
