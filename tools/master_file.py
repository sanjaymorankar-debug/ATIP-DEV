"""
The ATIP master file: docs/master/ATIP_MASTER_FILE_<date>.md and its workbook.

    python tools/master_file.py simulate   # re-run the exit-geometry simulations -> docs/master/simulations.json
    python tools/master_file.py diagrams   # render the .md's Mermaid diagrams -> docs/master/diagrams/*.png
    python tools/master_file.py build      # refresh the generated sections of the .md and write the .xlsx

`diagrams` needs the Mermaid CLI: `npm install @mermaid-js/mermaid-cli`, then MMDC=<path to mmdc> (and, where
Chromium is not Puppeteer's own, MMDC_PUPPETEER_CONFIG=<json with "executablePath">). `build` embeds the
PNGs it finds (needs Pillow) and otherwise keeps only the Mermaid source.

The markdown is the source. `build` reads its tables (terminology, layers, prompt review, objective,
roadmap) and the feature tracker (docs/ATIP_MASTER_TRACKER.csv), so nothing is typed twice:
  * between <!-- BEGIN:NAME --> / <!-- END:NAME --> (or a bare <!-- NAME --> on first build) the .md gets
    the generated FEATURE_SUMMARY and SIM_TABLES sections;
  * the workbook gets one sheet per topic, with every tracked feature plus one row per requirement of
    the consolidated development prompt (§6 of the .md, IDs CP-<phase>).

The simulations are synthetic geometric random walks with fixed seeds. They measure what a target / stop
geometry yields at a given underlying edge (annual drift); they say nothing about any ATIP signal. The
policy grid replays bars through strategy.backtest_aggressive.simulate(), the exit code live trading
uses. No database is opened.
"""

from __future__ import annotations

import csv
import json
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
MASTER_DIR = ROOT / "docs" / "master"
MD = MASTER_DIR / "ATIP_MASTER_FILE_2026-10-26.md"
XLSX = MD.with_suffix(".xlsx")
SIMS = MASTER_DIR / "simulations.json"
DIAGRAM_DIR = MASTER_DIR / "diagrams"
TRACKER = ROOT / "docs" / "ATIP_MASTER_TRACKER.csv"

COST_PCT = 0.25          # round trip, as strategy/backtest_history.DEFAULT_COST_PCT
TARGET_PCT = 4.0         # the owner's per-trade target


# ═══════════════════════════════════════════════════════════════════════════
#  SIMULATIONS
# ═══════════════════════════════════════════════════════════════════════════

def _paths(rng, n, sessions, steps, sigma_d, mu_a):
    """Price / entry for n paths, `steps` checks a session (log-normal, drift mu_a %/year)."""
    import numpy as np
    mu_d = np.log(1 + mu_a / 100) / 250
    s = sigma_d / 100 / np.sqrt(steps)
    m = mu_d / steps - 0.5 * s * s
    return np.exp(np.cumsum(rng.normal(m, s, size=(n, sessions * steps)), axis=1))


def _first_touch(path, target, stop):
    """Net % per path: +target or -stop at the first touch, else the last price; minus costs."""
    import numpy as np
    up, dn = path >= 1 + target / 100, path <= 1 - stop / 100
    iu = np.where(up.any(1), up.argmax(1), 10 ** 9)
    idn = np.where(dn.any(1), dn.argmax(1), 10 ** 9)
    hit_t, hit_s = iu < idn, idn < iu
    net = np.where(hit_t, target, np.where(hit_s, -stop, (path[:, -1] - 1) * 100)) - COST_PCT
    return hit_t, hit_s, net


def geometry_grid(n=20000, sessions=15, steps=25, seed=20261026):
    """Win rate and expectancy of a +4% target against each stop, at each volatility and edge."""
    import numpy as np
    rng = np.random.default_rng(seed)
    out = []
    for sigma_d in (1.5, 2.0, 2.5):
        for mu_a in (0.0, 20.0, 40.0):
            path = _paths(rng, n, sessions, steps, sigma_d, mu_a)
            for stop in (2.0, 3.0, 4.0, 6.0, 8.0, 12.0):
                hit_t, hit_s, net = _first_touch(path, TARGET_PCT, stop)
                out.append({"daily_vol_pct": sigma_d, "annual_drift_pct": mu_a, "target_pct": TARGET_PCT,
                            "stop_pct": stop, "horizon_sessions": sessions,
                            "p_target_first_pct": round(float(hit_t.mean()) * 100, 1),
                            "p_stop_first_pct": round(float(hit_s.mean()) * 100, 1),
                            "p_timeout_pct": round(float((~(hit_t | hit_s)).mean()) * 100, 1),
                            "win_rate_pct": round(float((net > 0).mean()) * 100, 1),
                            "expectancy_pct": round(float(net.mean()), 3),
                            "no_edge_formula_pct": round(stop / (TARGET_PCT + stop) * 100, 1)})
    return out


def required_drift(n=20000, sessions=15, steps=25, seed=1):
    """How much edge a 75% win rate at +4% needs, with a 3% and a 4% stop."""
    import numpy as np
    rng = np.random.default_rng(seed)
    out = []
    for sigma_d in (1.5, 2.0):
        for mu_a in (40.0, 80.0, 120.0, 200.0, 300.0):
            path = _paths(rng, n, sessions, steps, sigma_d, mu_a)
            for stop in (3.0, 4.0):
                hit_t, _, net = _first_touch(path, TARGET_PCT, stop)
                out.append({"daily_vol_pct": sigma_d, "annual_drift_pct": mu_a, "stop_pct": stop,
                            "p_target_first_pct": round(float(hit_t.mean()) * 100, 1),
                            "win_rate_pct": round(float((net > 0).mean()) * 100, 1),
                            "expectancy_pct": round(float(net.mean()), 2)})
    return out


# name: (aggressive-config overrides, initial stop %, sell everything at T1)
POLICIES = {
    "A: current default (T1 +3% sell 50%, T2 +6%, 3% stop)": ({}, 3.0, False),
    "B: flat +3% target, 3% stop": ({}, 3.0, True),
    "C: T1 +4% sell 50%, T2 +8%, 3% stop": ({"target_1_pct": 4.0, "target_2_pct": 8.0}, 3.0, False),
    "D: flat +4% target, 3% stop": ({"target_1_pct": 4.0, "target_2_pct": 8.0}, 3.0, True),
    "E: flat +4% target, 6% stop": ({"target_1_pct": 4.0, "target_2_pct": 8.0}, 6.0, True),
    "F: T1 +4% sell 70%, T2 +8%, 6% stop": ({"target_1_pct": 4.0, "target_1_exit_pct": 70.0,
                                              "target_2_pct": 8.0}, 6.0, False),
}


def policy_grid(n=4000, hold=30, pre=14, steps=25, seed=7):
    """Each exit policy through strategy.backtest_aggressive.simulate() on the same synthetic bars."""
    import numpy as np
    sys.path.insert(0, str(ROOT))
    from strategy.aggressive import DEFAULTS, validate_config
    from strategy.backtest_aggressive import Bar, metrics, simulate

    rng = np.random.default_rng(seed)
    out = []
    for sigma_d in (1.5, 2.0):
        for mu_a in (0.0, 20.0, 40.0):
            lp = np.log(_paths(rng, n, pre + hold, steps, sigma_d, mu_a)).reshape(n, pre + hold, steps)
            hi, lo, cl = lp.max(2), lp.min(2), lp[:, :, -1]
            entry = cl[:, pre - 1:pre]                    # enter at the close before the window
            hi, lo, cl = (np.exp(a - entry) * 100 for a in (hi, lo, cl))
            atrs = []
            for k in range(n):                            # ATR% from the pre-entry bars only
                tr = [hi[k, 0] - lo[k, 0]] + [max(hi[k, j] - lo[k, j], abs(hi[k, j] - cl[k, j - 1]),
                                                  abs(lo[k, j] - cl[k, j - 1])) for j in range(1, pre)]
                atrs.append(float(np.mean(tr) / cl[k, pre - 1] * 100))
            for name, (over, stop, flat) in POLICIES.items():
                cfg = validate_config({**DEFAULTS, **over})
                trades = [simulate(100.0, 100,
                                   [Bar(high=float(hi[k, j]), low=float(lo[k, j]), close=float(cl[k, j]))
                                    for j in range(pre, pre + hold)],
                                   None, cfg, stop_pct=stop, cost_pct=COST_PCT, exit_all_at_t1=flat,
                                   atr_pct=atrs[k])
                          for k in range(n)]
                m = metrics(trades)
                out.append({"daily_vol_pct": sigma_d, "annual_drift_pct": mu_a, "policy": name, "stop_pct": stop,
                            **{k: m[k] for k in ("win_rate_pct", "expectancy_pct", "profit_factor",
                                                 "target_1_hit_rate_pct", "initial_stop_exit_pct",
                                                 "trailing_exit_pct", "avg_bars_held")}})
    return out


def simulate_all() -> dict:
    res = {"note": "Synthetic geometric random walks, fixed seeds; costs 0.25% round trip. "
                   "Measures exit geometry at a given edge (drift), not any ATIP signal.",
           "geometry": geometry_grid(), "required_drift": required_drift(), "policies": policy_grid()}
    SIMS.write_text(json.dumps(res, indent=1), encoding="utf-8")
    return res


# ═══════════════════════════════════════════════════════════════════════════
#  MARKDOWN
# ═══════════════════════════════════════════════════════════════════════════

def md_tables(text: str, heading: str) -> list[dict]:
    """
    Every pipe table under the heading that starts with `heading`, up to the next heading of the same or a
    higher level: [{"section": nearest sub-heading, "header": [...], "rows": [[...], ...]}].
    """
    lines = text.splitlines()
    start = next((i for i, l in enumerate(lines) if l.startswith("#") and l.lstrip("#").strip().startswith(heading)),
                 None)
    if start is None:
        raise LookupError(f"heading {heading!r} not found in {MD.name}")
    level = len(lines[start]) - len(lines[start].lstrip("#"))
    section, out, cur, fence = lines[start].lstrip("#").strip(), [], None, False
    for l in lines[start + 1:]:
        if l.startswith("```"):
            fence = not fence
        if fence:
            continue
        if l.startswith("#"):
            if len(l) - len(l.lstrip("#")) <= level:
                break
            section = l.lstrip("#").strip()
        if l.startswith("|"):
            cells = [c.strip() for c in l.strip().strip("|").split("|")]
            if cur is None:
                cur = {"section": section, "header": cells, "rows": []}
                out.append(cur)
            elif not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                cur["rows"].append(cells)
        else:
            cur = None
    return out


def _plain(s: str) -> str:
    """Markdown cell -> plain text for the workbook."""
    s = re.sub(r"\*\*(.+?)\*\*", r"\1", s)
    s = re.sub(r"(?<!\w)\*(.+?)\*(?!\w)", r"\1", s)
    s = re.sub(r"`(.+?)`", r"\1", s)
    s = re.sub(r"\[(.+?)\]\(.+?\)", r"\1", s)
    return s.replace("\\|", "|").strip()


def _replace_block(text: str, name: str, body: str) -> str:
    block = f"<!-- BEGIN:{name} -->\n{body.rstrip()}\n<!-- END:{name} -->"
    pat = re.compile(rf"<!-- BEGIN:{name} -->.*?<!-- END:{name} -->", re.DOTALL)
    if pat.search(text):
        return pat.sub(lambda _: block, text)
    if f"<!-- {name} -->" not in text:
        raise LookupError(f"no <!-- {name} --> marker in {MD.name}")
    return text.replace(f"<!-- {name} -->", block)


def mermaid_blocks(text: str) -> list[tuple[str, str]]:
    """[(heading just above the block, Mermaid source)] for every ```mermaid block, in order."""
    return re.findall(r"### ([^\n]+)\n\n```mermaid\n(.*?)```", text, re.DOTALL)


def _slug(name: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-")


def render_diagrams() -> list[Path]:
    """Each Mermaid block -> DIAGRAM_DIR/<slug>.png through the Mermaid CLI (see the module docstring)."""
    import os
    import shutil
    import subprocess
    import tempfile
    mmdc = os.environ.get("MMDC") or shutil.which("mmdc")
    if not mmdc:
        raise SystemExit("mmdc not found: install @mermaid-js/mermaid-cli and set MMDC=<path to mmdc>")
    DIAGRAM_DIR.mkdir(parents=True, exist_ok=True)
    out = []
    with tempfile.TemporaryDirectory() as tmp:
        for name, src in mermaid_blocks(MD.read_text(encoding="utf-8")):
            mmd = Path(tmp) / f"{_slug(name)}.mmd"
            mmd.write_text(src, encoding="utf-8")
            png = DIAGRAM_DIR / f"{_slug(name)}.png"
            cmd = [mmdc, "-i", str(mmd), "-o", str(png), "-w", "1800", "-b", "white"]
            if os.environ.get("MMDC_PUPPETEER_CONFIG"):
                cmd += ["-p", os.environ["MMDC_PUPPETEER_CONFIG"]]
            subprocess.run(cmd, check=True, capture_output=True)
            out.append(png)
    return out


def load_tracker() -> list[dict]:
    with TRACKER.open(encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


STATUS_ORDER = ("COMPLETED", "IMPLEMENTED BUT NOT VERIFIED", "IN PROGRESS", "BLOCKED")


def feature_summary_md(rows: list[dict]) -> str:
    by = defaultdict(Counter)
    for r in rows:
        by[r["Module"]][r["Status"]] += 1
    head = "| Module | Features | " + " | ".join(s.title() for s in STATUS_ORDER) + " |"
    out = [head, "|" + "---|" * (len(STATUS_ORDER) + 2)]
    for mod, c in sorted(by.items(), key=lambda kv: -sum(kv[1].values())):
        out.append(f"| {mod} | {sum(c.values())} | " + " | ".join(str(c.get(s, 0)) for s in STATUS_ORDER) + " |")
    tot = Counter(r["Status"] for r in rows)
    out.append(f"| **All** | **{len(rows)}** | " + " | ".join(f"**{tot.get(s, 0)}**" for s in STATUS_ORDER) + " |")
    return "\n".join(out)


def sim_tables_md(sims: dict) -> str:
    g = sims["geometry"]
    out = ["**Table A: a +4% target against each stop, 2.0% daily volatility, 15 sessions.** Win rate / net",
           "expectancy per trade, by edge.", "",
           "| Stop | No-edge formula | 0% drift: win / E | 20% drift: win / E | 40% drift: win / E |",
           "|---|---|---|---|---|"]
    for stop in (2.0, 3.0, 4.0, 6.0, 8.0, 12.0):
        cells = []
        for mu in (0.0, 20.0, 40.0):
            r = next(x for x in g if x["daily_vol_pct"] == 2.0 and x["annual_drift_pct"] == mu and x["stop_pct"] == stop)
            cells.append(f"{r['win_rate_pct']:.1f}% / {r['expectancy_pct']:+.2f}%")
        out.append(f"| {stop:g}% | {stop / (TARGET_PCT + stop) * 100:.0f}% | " + " | ".join(cells) + " |")
    best = max(g, key=lambda x: x["win_rate_pct"])
    out += ["", (f"Across all {len(g)} cells (1.5–2.5% volatility, 0–40% drift), the highest win rate is "
                f"{best['win_rate_pct']:.1f}%, with a {best['stop_pct']:g}% stop at "
                f"{best['daily_vol_pct']:g}% volatility and {best['annual_drift_pct']:g}% drift; its expectancy is "
                f"{best['expectancy_pct']:+.2f}%. The full grid is on the workbook's Sim_Geometry sheet."), ""]

    p = sims["policies"]
    out += ["**Table B: exit policies through ATIP's own `simulate()` (30-session window, ATR trail).** Win rate /",
            "expectancy per trade.", "",
            "| Policy | Vol | 0% drift | 20% drift | 40% drift |", "|---|---|---|---|---|"]
    for name in POLICIES:
        for vol in (1.5, 2.0):
            cells = []
            for mu in (0.0, 20.0, 40.0):
                r = next(x for x in p if x["policy"] == name and x["daily_vol_pct"] == vol and x["annual_drift_pct"] == mu)
                cells.append(f"{r['win_rate_pct']:.1f}% / {r['expectancy_pct']:+.2f}%")
            out.append(f"| {name} | {vol:g}% | " + " | ".join(cells) + " |")
    out.append("")

    d = sims["required_drift"]
    out += ["**Table C: the edge that a 75% win rate at +4% needs (15 sessions).** Win rate / expectancy.", "",
            "| Annual drift | 1.5% vol, 3% stop | 1.5% vol, 4% stop | 2.0% vol, 3% stop | 2.0% vol, 4% stop |",
            "|---|---|---|---|---|"]
    for mu in sorted({x["annual_drift_pct"] for x in d}):
        cells = []
        for vol, stop in ((1.5, 3.0), (1.5, 4.0), (2.0, 3.0), (2.0, 4.0)):
            r = next(x for x in d if x["annual_drift_pct"] == mu and x["daily_vol_pct"] == vol and x["stop_pct"] == stop)
            mark = " ✓" if r["win_rate_pct"] >= 75 else ""
            cells.append(f"{r['win_rate_pct']:.1f}% / {r['expectancy_pct']:+.2f}%{mark}")
        out.append(f"| {mu:g}% | " + " | ".join(cells) + " |")
    out += ["", "✓ marks the cells that reach 75%."]
    return "\n".join(out)


# ═══════════════════════════════════════════════════════════════════════════
#  WORKBOOK
# ═══════════════════════════════════════════════════════════════════════════

FEATURE_COLS = ("ID", "Phase Name", "Module", "Feature", "Target Requirement", "Priority", "Status",
                "Testing Status", "Completion %", "Current Implementation", "Next Action / Missing Work",
                "Key Files", "Deployment Wave", "Blocker / Input Required")
STATUS_FILL = {"COMPLETED": "C6EFCE", "PRESENT": "C6EFCE", "IMPLEMENTED BUT NOT VERIFIED": "DDEBF7",
               "PARTIAL": "FFEB9C", "IN PROGRESS": "FFEB9C", "BLOCKED": "F8CBAD", "ABSENT": "F8CBAD",
               "DEFECT": "F8CBAD", "NEW": "EDEDED"}


def _sheet(wb, title, header, rows, widths=None, status_col=None):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    ws = wb.create_sheet(title[:31])
    ws.append(list(header))
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3A5F")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        ws.append(list(r))
    si = header.index(status_col) + 1 if status_col in header else None
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        if si:
            v = str(row[si - 1].value or "").split()[0].upper() if row[si - 1].value else ""
            v = next((k for k in STATUS_FILL if str(row[si - 1].value or "").upper().startswith(k)), v)
            if v in STATUS_FILL:
                row[si - 1].fill = PatternFill("solid", fgColor=STATUS_FILL[v])
    for i, w in enumerate(widths or [], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions
    return ws


def _diagram_sheets(wb, diagrams):
    """One sheet per rendered diagram, the picture scaled to 1,100 px wide. Skipped without Pillow."""
    try:
        import PIL  # noqa: F401  (openpyxl needs it to read the PNG)
        from openpyxl.drawing.image import Image
    except ImportError:
        return
    from openpyxl.styles import Font
    for i, (name, _src) in enumerate(diagrams, 1):
        png = DIAGRAM_DIR / f"{_slug(name)}.png"
        if not png.exists():
            continue
        ws = wb.create_sheet(f"Diagram_{i}")
        ws["A1"] = name
        ws["A1"].font = Font(bold=True, size=14)
        img = Image(str(png))
        scale = min(1.0, 1100 / img.width)
        img.width, img.height = int(img.width * scale), int(img.height * scale)
        ws.add_image(img, "A3")


def _flatten(tables, extra_first=None):
    """
    Tables joined into one sheet, each cell placed under its own column name (tables under one heading need
    not share columns), with the section name in front when extra_first is set.
    """
    header = [extra_first] if extra_first else []
    for t in tables:
        header += [h for h in t["header"] if h not in header]
    rows = []
    for t in tables:
        d = [dict(zip(t["header"], (_plain(c) for c in r))) for r in t["rows"]]
        rows += [[t["section"] if extra_first and h == extra_first else x.get(h, "") for h in header] for x in d]
    return header, rows


def prompt_rows(text: str) -> list[list]:
    """§6's review table(s) as feature rows CP-<n>, so the new requirements sit in the feature list."""
    out, seen = [], Counter()
    for t in md_tables(text, "6. Review"):
        h = [c.lower() for c in t["header"]]
        if "status" not in h or "requirement" not in h:
            continue
        for r in t["rows"]:
            d = dict(zip(h, (_plain(c) for c in r)))
            phase = int(re.sub(r"\D", "", d.get("phase", "")) or 0)
            seen[phase] += 1
            out.append({"ID": f"CP-{phase:02d}-{seen[phase]:02d}", "Phase Name": "Consolidated prompt (2026-10-26)",
                        "Module": d.get("area") or d.get("phase", ""), "Feature": d.get("requirement", ""),
                        "Target Requirement": d.get("requirement", ""), "Priority": d.get("priority", ""),
                        "Status": d.get("status", "").upper(), "Testing Status": "",
                        "Completion %": "", "Current Implementation": d.get("evidence", ""),
                        "Next Action / Missing Work": d.get("gap / action", d.get("gap", "")),
                        "Key Files": "", "Deployment Wave": "", "Blocker / Input Required": ""})
    return out


def build_xlsx(text: str, tracker: list[dict], sims: dict):
    from openpyxl import Workbook
    wb = Workbook()
    wb.remove(wb.active)

    readme = [
        ("Title", "ATIP master file: architecture, process flows, terminology, features, 4% / 75% objective"),
        ("Source", (f"docs/master/{MD.name} (the markdown is the source; this workbook is generated by "
                   "tools/master_file.py build)")),
        ("Features", (f"{len(tracker)} tracked features from docs/ATIP_MASTER_TRACKER.csv plus the consolidated "
                     "prompt's requirements (IDs CP-...)")),
        ("Simulations", sims.get("note", "")),
        ("Diagrams", "Mermaid source on the Diagrams sheet; the .md renders them on GitHub"),
    ]
    _sheet(wb, "README", ["Item", "Value"], readme, [18, 120])

    summary = re.search(r"## 1\. Executive summary\n(.*?)\n---", text, re.DOTALL)
    paras = [p.strip() for p in (summary.group(1) if summary else "").split("\n\n") if p.strip()]
    _sheet(wb, "Executive_Summary", ["Paragraph"], [[_plain(" ".join(p.split()))] for p in paras], [160])

    h, rows = _flatten(md_tables(text, "2.2 Layers"))
    _sheet(wb, "Architecture_Layers", h, rows, [26, 50, 24, 40])

    diagrams = mermaid_blocks(text)
    _sheet(wb, "Diagrams", ["Diagram", "Mermaid source", "Picture"],
           [[n, src, f"sheet Diagram_{i}" if (DIAGRAM_DIR / f"{_slug(n)}.png").exists() else "not rendered"]
            for i, (n, src) in enumerate(diagrams, 1)], [40, 120, 16])
    _diagram_sheets(wb, diagrams)

    h, rows = _flatten(md_tables(text, "4. Terminology"), extra_first="Group")
    _sheet(wb, "Terminology", h, rows, [30, 26, 24, 90, 40])

    cp = prompt_rows(text)
    feats = [[r.get(c, "") for c in FEATURE_COLS] for r in tracker] + [[r[c] for c in FEATURE_COLS] for r in cp]
    _sheet(wb, "Features", list(FEATURE_COLS), feats, [12, 22, 18, 36, 40, 8, 18, 14, 8, 50, 50, 36, 8, 30],
           status_col="Status")

    by = defaultdict(Counter)
    for r in tracker:
        by[r["Module"]][r["Status"]] += 1
    srows = [[m, sum(c.values())] + [c.get(s, 0) for s in STATUS_ORDER]
             for m, c in sorted(by.items(), key=lambda kv: -sum(kv[1].values()))]
    _sheet(wb, "Feature_Summary", ["Module", "Features", *STATUS_ORDER], srows, [28, 10, 14, 30, 14, 10])

    h, rows = _flatten(md_tables(text, "6. Review"))
    _sheet(wb, "Prompt_Review", h, rows, [10, 22, 40, 12, 60, 60], status_col="Status")

    obj = md_tables(text, "7. The objective")
    h, rows = _flatten([t for t in obj if t["header"][0] in ("Stop S",)])
    _sheet(wb, "Objective_Arithmetic", h, rows, [10, 30, 16, 22])
    h, rows = _flatten([t for t in obj if t["header"][0] in ("Target as stated",)])
    _sheet(wb, "Objective_Verdict", h, rows, [60, 30, 70])
    h, rows = _flatten([t for t in obj if t["header"][:2] == ["#", "Lever"]])
    _sheet(wb, "Levers", h, rows, [4, 30, 70, 50, 16])
    h, rows = _flatten([t for t in obj if t["header"][0] == "Stage"])
    _sheet(wb, "Validation_Gates", h, rows, [26, 80, 40])
    h, rows = _flatten([t for t in obj if t["header"][0] == "Evidence"])
    _sheet(wb, "Measured_Evidence", h, rows, [40, 70, 50])

    g = sims["geometry"]
    _sheet(wb, "Sim_Geometry", list(g[0]), [list(r.values()) for r in g])
    p = sims["policies"]
    _sheet(wb, "Sim_Policies", list(p[0]), [list(r.values()) for r in p], [8, 8, 52])
    d = sims["required_drift"]
    _sheet(wb, "Sim_Required_Edge", list(d[0]), [list(r.values()) for r in d])

    h, rows = _flatten(md_tables(text, "8. Roadmap"))
    _sheet(wb, "Roadmap", h, rows, [5, 8, 80, 60, 10])

    wb.save(XLSX)
    return {"sheets": wb.sheetnames, "features": len(feats), "prompt_rows": len(cp)}


def build():
    if not SIMS.exists():
        simulate_all()
    sims = json.loads(SIMS.read_text(encoding="utf-8"))
    tracker = load_tracker()
    text = MD.read_text(encoding="utf-8")
    text = _replace_block(text, "FEATURE_SUMMARY", feature_summary_md(tracker))
    text = _replace_block(text, "SIM_TABLES", sim_tables_md(sims))
    MD.write_text(text, encoding="utf-8")
    return build_xlsx(text, tracker, sims)


if __name__ == "__main__":
    cmd = sys.argv[1] if len(sys.argv) > 1 else "build"
    if cmd == "simulate":
        r = simulate_all()
        print(f"{SIMS.relative_to(ROOT)}: {len(r['geometry'])} geometry, {len(r['policies'])} policy, "
              f"{len(r['required_drift'])} required-edge cells")
    elif cmd == "diagrams":
        for p in render_diagrams():
            print(p.relative_to(ROOT))
    elif cmd == "build":
        print(build())
    else:
        sys.exit(__doc__)
