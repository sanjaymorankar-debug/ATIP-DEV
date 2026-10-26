"""
tools/master_file.py: the master file's generator reads its own markdown, so these check that the parts it
depends on hold -- the tables it lifts into the workbook, the generated blocks, the CP-<phase> ids -- and
that the exit-geometry simulation agrees with the closed form it is quoted next to.
"""

import importlib.util
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture(scope="module")
def mf():
    spec = importlib.util.spec_from_file_location("master_file", ROOT / "tools" / "master_file.py")
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    return m


@pytest.fixture(scope="module")
def text(mf):
    return mf.MD.read_text(encoding="utf-8")


def test_every_review_row_has_a_known_status_and_evidence(mf, text):
    (table,) = [t for t in mf.md_tables(text, "6. Review") if "Status" in t["header"]]
    h = table["header"]
    assert h == ["Phase", "Area", "Requirement", "Status", "Evidence", "Gap / action"]
    assert len(table["rows"]) >= 30
    for r in table["rows"]:
        assert len(r) == len(h), r
        assert r[3] in ("PRESENT", "PARTIAL", "ABSENT", "DEFECT"), r
        assert r[0].startswith("P") and r[4], r


def test_prompt_rows_are_numbered_within_their_phase(mf, text):
    rows = mf.prompt_rows(text)
    ids = [r["ID"] for r in rows]
    assert len(ids) == len(set(ids))
    assert ids[0] == "CP-01-01"
    p3 = [i for i in ids if i.startswith("CP-03-")]
    assert p3 == [f"CP-03-{n:02d}" for n in range(1, len(p3) + 1)]


def test_tables_join_by_column_name(mf):
    md = ("## 4. Terminology\n\n### 4.1 A\n\n| Term | Full name | Meaning in ATIP | Where |\n|---|---|---|---|\n"
          "| VPI | Volatility-Potential Index | swing room | engine |\n\n### 4.5 B\n\n| Term | Meaning in ATIP |\n"
          "|---|---|\n| NSE | exchange |\n\n## 5. Next\n\n| X |\n|---|\n| not mine |\n")
    header, rows = mf._flatten(mf.md_tables(md, "4. Terminology"), extra_first="Group")
    assert header == ["Group", "Term", "Full name", "Meaning in ATIP", "Where"]
    assert rows == [["4.1 A", "VPI", "Volatility-Potential Index", "swing room", "engine"],
                    ["4.5 B", "NSE", "", "exchange", ""]]


def test_generated_blocks_replace_cleanly_and_idempotently(mf):
    once = mf._replace_block("a\n<!-- X -->\nb", "X", "| t |")
    assert once == "a\n<!-- BEGIN:X -->\n| t |\n<!-- END:X -->\nb"
    assert mf._replace_block(once, "X", "| t |") == once
    assert mf._replace_block(once, "X", "| u |").count("| u |") == 1
    with pytest.raises(LookupError):
        mf._replace_block("no marker", "X", "body")


def test_no_edge_win_rate_matches_the_closed_form(mf):
    """With no drift, P(+T before -S) ~ S / (T + S): the 'no-edge formula' column the report quotes."""
    pytest.importorskip("numpy")
    grid = mf.geometry_grid(n=4000, sessions=60, steps=25, seed=3)
    for r in grid:
        if r["annual_drift_pct"] == 0 and r["daily_vol_pct"] == 2.5 and r["stop_pct"] <= 6:
            assert r["p_timeout_pct"] < 5, r
            assert abs(r["p_target_first_pct"] - r["no_edge_formula_pct"]) < 4, r
            assert r["expectancy_pct"] < 0, r          # no edge: costs make it a loser


def test_policy_grid_runs_through_the_live_exit_code(mf):
    pytest.importorskip("numpy")
    out = mf.policy_grid(n=60, hold=10, seed=5)
    assert {r["policy"] for r in out} == set(mf.POLICIES)
    for r in out:
        assert 0 <= r["win_rate_pct"] <= 100 and r["avg_bars_held"] <= 10


def test_workbook_carries_every_tracked_feature(mf, text, tmp_path, monkeypatch):
    pytest.importorskip("openpyxl")
    from openpyxl import load_workbook
    monkeypatch.setattr(mf, "XLSX", tmp_path / "master.xlsx")
    sims = {"note": "", "geometry": mf.geometry_grid(n=200), "required_drift": mf.required_drift(n=200),
            "policies": mf.policy_grid(n=20, hold=10)}
    tracker = mf.load_tracker()
    r = mf.build_xlsx(text, tracker, sims)
    assert r["features"] == len(tracker) + r["prompt_rows"]
    wb = load_workbook(tmp_path / "master.xlsx", read_only=True)
    ids = [row[0] for row in wb["Features"].iter_rows(min_row=2, values_only=True)]
    assert {t["ID"] for t in tracker} <= set(ids)
    assert wb["Terminology"].max_row > 100
